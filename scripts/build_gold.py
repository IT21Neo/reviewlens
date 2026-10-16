"""Milestone 4b — validate the hand labels and write data/eval/gold_labels.jsonl.

Reads your labels from data/eval/labeling.xlsx (made by scripts/make_label_sheet.py), or
from data/eval/to_label.csv if no workbook exists, and checks every row:
  - review_id exists in data/reviews.csv and its text was not changed,
  - gold_sentiment is one of app.taxonomy.SENTIMENTS,
  - every aspect in gold_aspects is in app.taxonomy.ASPECTS (no duplicates),
  - every evidence phrase appears verbatim in the review text.
If anything is wrong it prints ALL problems and writes nothing (fail closed).

Run from the repo root:   python -m scripts.build_gold
"""
from __future__ import annotations

import csv
import json
import sys
from pathlib import Path

from app.taxonomy import ASPECTS, SENTIMENTS

REVIEWS = Path("data/reviews.csv")
TO_LABEL = Path("data/eval/to_label.csv")
XLSX = Path("data/eval/labeling.xlsx")
GOLD = Path("data/eval/gold_labels.jsonl")
# human = typed by the annotator; claude_checked / claude_edited = AI draft reviewed by the annotator
LABEL_SOURCES = {"human", "claude_checked", "claude_edited"}


def parse_list(cell: str, sep: str) -> list[str]:
    return [x.strip() for x in (cell or "").split(sep) if x.strip()]


def validate(label_rows: list[dict], reviews: dict[str, str]) -> tuple[list[dict], list[str]]:
    gold, errors, seen = [], [], set()
    for line, row in enumerate(label_rows, start=2):  # line 1 is the header
        rid = (row.get("review_id") or "").strip()
        where = f"line {line} ({rid or 'no id'})"
        if rid not in reviews:
            errors.append(f"{where}: review_id not found in {REVIEWS}")
            continue
        if rid in seen:
            errors.append(f"{where}: duplicate review_id")
        seen.add(rid)
        text = reviews[rid]
        if (row.get("review_text") or "").replace("\r\n", "\n").replace("_x000D_", "") != text:
            errors.append(f"{where}: review_text was edited — restore the original text")

        sentiment = (row.get("gold_sentiment") or "").strip().lower()
        if sentiment not in SENTIMENTS:
            errors.append(f"{where}: gold_sentiment {sentiment!r} must be one of {list(SENTIMENTS)}")

        aspects = [a.lower() for a in parse_list(row.get("gold_aspects", ""), ";")]
        bad = [a for a in aspects if a not in ASPECTS]
        if bad:
            errors.append(f"{where}: unknown aspect(s) {bad}; allowed {list(ASPECTS)}")
        if len(set(aspects)) != len(aspects):
            errors.append(f"{where}: duplicate aspect in {aspects}")

        evidence = parse_list(row.get("gold_evidence", ""), "|")
        missing = [e for e in evidence if e not in text]
        if missing:
            errors.append(f"{where}: evidence not found verbatim in the review: {missing}")

        source = (row.get("label_source") or "").strip() or "human"
        if source == "claude_draft":
            errors.append(f"{where}: AI draft not reviewed yet — read the review, then set label_source "
                          "to claude_checked (kept) or claude_edited (changed)")
        elif source not in LABEL_SOURCES:
            errors.append(f"{where}: label_source {source!r} must be one of {sorted(LABEL_SOURCES)}")

        ordered = [a for a in ASPECTS if a in aspects]  # canonical order, duplicates removed
        gold.append({"review_id": rid, "split": row.get("split", "").strip(),
                     "gold_sentiment": sentiment, "gold_aspects": ordered,
                     "gold_evidence": evidence, "notes": (row.get("notes") or "").strip(),
                     "label_source": source})
    return gold, errors


def read_label_rows() -> tuple[list[dict], Path]:
    if XLSX.exists():
        from openpyxl import load_workbook

        ws = load_workbook(XLSX, read_only=True, data_only=True)["labels"]
        it = ws.iter_rows(values_only=True)
        header = [str(h) for h in next(it)]
        rows = [{h: ("" if v is None else str(v)) for h, v in zip(header, values)}
                for values in it if any(v is not None for v in values)]
        return rows, XLSX
    with open(TO_LABEL, encoding="utf-8-sig", newline="") as f:
        return list(csv.DictReader(f)), TO_LABEL


def main() -> int:
    with open(REVIEWS, encoding="utf-8", newline="") as f:
        reviews = {r["review_id"]: r["review_text"] for r in csv.DictReader(f)}
    rows, source = read_label_rows()
    print(f"reading labels from {source}")

    gold, errors = validate(rows, reviews)
    if errors:
        print(f"{len(errors)} problem(s) — nothing written:")
        for e in errors:
            print("  -", e)
        return 1

    with open(GOLD, "w", encoding="utf-8", newline="\n") as f:
        for g in gold:
            f.write(json.dumps(g, ensure_ascii=False) + "\n")
    splits = {s: sum(g["split"] == s for g in gold) for s in ("dev", "test")}
    sent = {s: sum(g["gold_sentiment"] == s for g in gold) for s in SENTIMENTS}
    asp = {a: sum(a in g["gold_aspects"] for g in gold) for a in ASPECTS}
    print(f"wrote {len(gold)} gold labels to {GOLD}")
    print("splits    :", splits)
    print("sentiment :", sent)
    print("aspects   :", asp)
    src = {s: sum(g["label_source"] == s for g in gold) for s in sorted(LABEL_SOURCES)}
    print("source    :", src)
    return 0


if __name__ == "__main__":
    sys.exit(main())
