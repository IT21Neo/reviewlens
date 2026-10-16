"""Milestone 4a (2) — turn data/eval/to_label.csv into an Excel workbook that is easy to label.

Why: opening a CSV with long, multi-line Thai text directly in Excel can put everything in
one column. A real .xlsx keeps each field in its own column and adds:
  - wide, wrapped text columns and a frozen header row,
  - a drop-down list for gold_sentiment (positive / neutral / negative),
  - a second sheet "guide" with the allowed aspects.

Run from the repo root:   python -m scripts.make_label_sheet
Then open data/eval/labeling.xlsx in Excel, fill in the 4 empty columns, and save it
(normal Ctrl+S, keep the .xlsx format). `python -m scripts.build_gold` reads this file.
Refuses to overwrite an existing workbook (your labels) unless you pass --force.
`--guide-only` rewrites just the Thai "guide" sheet of an existing workbook.
"""
from __future__ import annotations

import csv
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

from app.taxonomy import SENTIMENTS
from scripts.label_guide_th import ASPECTS_TH, STEPS, TRICKY

TO_LABEL = Path("data/eval/to_label.csv")
XLSX = Path("data/eval/labeling.xlsx")
COLUMNS = ["review_id", "split", "review_text", "gold_sentiment", "gold_aspects", "gold_evidence", "notes"]
WIDTHS = {"A": 10, "B": 7, "C": 90, "D": 16, "E": 28, "F": 50, "G": 30}


def build(rows: list[dict], path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "labels"
    ws.append(COLUMNS)
    for r in rows:
        ws.append([r.get(c, "") for c in COLUMNS])

    head_fill = PatternFill("solid", fgColor="DDEBF7")
    todo_fill = PatternFill("solid", fgColor="FFF2CC")
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = head_fill
    for col, width in WIDTHS.items():
        ws.column_dimensions[col].width = width
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            cell.data_type = "s"  # always plain text, never a formula
        for cell in row[3:]:
            cell.fill = todo_fill  # the columns you fill in
    ws.freeze_panes = "D2"

    dv = DataValidation(type="list", formula1='"' + ",".join(SENTIMENTS) + '"', allow_blank=True,
                        showErrorMessage=True, errorTitle="gold_sentiment",
                        error="Choose positive, neutral or negative")
    ws.add_data_validation(dv)
    dv.add(f"D2:D{len(rows) + 1}")

    write_guide(wb)
    wb.save(path)


def write_guide(wb) -> None:
    """(Re)create the Thai "guide" sheet. Never touches the "labels" sheet."""
    if "guide" in wb.sheetnames:
        del wb["guide"]
    g = wb.create_sheet("guide")
    bold, wrap = Font(bold=True), Alignment(wrap_text=True, vertical="top")
    head_fill = PatternFill("solid", fgColor="DDEBF7")

    def line(text):  # one sentence spread over the whole table width
        g.append([text])
        g.merge_cells(start_row=g.max_row, start_column=1, end_row=g.max_row, end_column=5)

    def title(text):
        line(text)
        g.cell(g.max_row, 1).font = Font(bold=True, size=13)

    def header(cols):
        g.append(cols)
        for c in g[g.max_row]:
            c.font, c.fill = bold, head_fill

    title("วิธีเลือก gold_aspects (กฎเต็ม: docs/annotation_guideline.md)")
    for s in STEPS:
        line(s)
    g.append([])

    title("แต่ละหมวดหมายถึงอะไร")
    header(["หมวด (พิมพ์ชื่อนี้)", "คำถามช่วยตัดสิน", "ใส่หมวดนี้เมื่อมีความเห็นเรื่อง…", "ตัวอย่างที่ใส่", "ไม่ใส่ / ไปหมวดอื่น"])
    for row in ASPECTS_TH:
        g.append(list(row))
        g.cell(g.max_row, 1).font = bold
    g.append([])

    title("ประโยคที่มักสับสน")
    header(["ประโยคในรีวิว", "ใส่ gold_aspects", "เหตุผล"])
    for row in TRICKY:
        g.append(list(row))
    g.append([])

    title("รูปแบบการกรอก")
    g.append(["gold_aspects", "พิมพ์ชื่อหมวดภาษาอังกฤษ คั่นด้วย ;   เช่น  taste;price"])
    g.append(["gold_evidence", "ไม่บังคับ — copy วลีจากรีวิวมาตรง ๆ คั่นด้วย  |  "])

    for col, width in {"A": 34, "B": 34, "C": 48, "D": 52, "E": 52}.items():
        g.column_dimensions[col].width = width
    for row in g.iter_rows():
        for c in row:
            c.alignment = wrap


def refresh_guide(path: Path) -> None:
    """Replace only the guide sheet of an existing workbook; labels you typed are kept."""
    from openpyxl import load_workbook

    wb = load_workbook(path)
    write_guide(wb)
    wb.save(path)


def main(argv: list[str]) -> int:
    if "--guide-only" in argv:
        refresh_guide(XLSX)
        print(f"updated the 'guide' sheet in {XLSX} (your labels were not changed)")
        return 0
    if XLSX.exists() and "--force" not in argv:
        sys.exit(f"{XLSX} already exists — not overwriting your labels. Use --force to replace it.")
    with open(TO_LABEL, encoding="utf-8-sig", newline="") as f:
        rows = list(csv.DictReader(f))
    build(rows, XLSX)
    print(f"wrote {len(rows)} rows to {XLSX} — open it in Excel and fill the yellow columns")
    return 0


if __name__ == "__main__":
    sys.exit(main(sys.argv[1:]))
