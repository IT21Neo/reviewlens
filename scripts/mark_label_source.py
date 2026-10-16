"""Record where each gold label came from, by comparing it with Claude's suggestions.

Use this when labels were typed into data/eval/labeling.xlsx while looking at Claude's
suggestions (data/eval/suggestions_dev_claude.csv and suggestions_test_claude.csv).
For every row whose label_source cell is still empty it writes:

  claude_checked — your sentiment AND aspect set are identical to Claude's suggestion
  claude_edited  — you differ from the suggestion in sentiment or aspects

Rows you already marked are left alone. If you labelled a row completely on your own,
type `human` in its label_source cell afterwards. Prints every difference so you can
re-check them, and a summary for the report.

Run from the repo root (Excel closed):   python -m scripts.mark_label_source
"""
from __future__ import annotations

import csv
import sys
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Font
from openpyxl.worksheet.datavalidation import DataValidation

XLSX = Path("data/eval/labeling.xlsx")
SUGGESTION_FILES = [Path("data/eval/suggestions_dev_claude.csv"), Path("data/eval/suggestions_test_claude.csv")]
CHOICES = ("human", "claude_checked", "claude_edited")


def aspect_set(cell) -> frozenset[str]:
    return frozenset(a.strip().lower() for a in str(cell or "").split(";") if a.strip())


def load_suggestions(files=SUGGESTION_FILES) -> dict[str, dict]:
    out = {}
    for path in files:
        with open(path, encoding="utf-8-sig", newline="") as f:
            for r in csv.DictReader(f):
                out[r["review_id"]] = r
    return out


def mark(xlsx: Path = XLSX, suggestions: dict | None = None) -> dict:
    suggestions = suggestions if suggestions is not None else load_suggestions()
    wb = load_workbook(xlsx)
    ws = wb["labels"]
    header = [c.value for c in ws[1]]
    if "label_source" not in header:
        ws.cell(1, len(header) + 1, "label_source").font = Font(bold=True)
        header.append("label_source")
    col = {name: i + 1 for i, name in enumerate(header)}

    counts = {c: 0 for c in CHOICES} | {"kept": 0, "no_suggestion": 0}
    diffs = []
    for r in range(2, ws.max_row + 1):
        rid = ws.cell(r, col["review_id"]).value
        if not rid:
            continue
        src = ws.cell(r, col["label_source"])
        if src.value:
            counts["kept"] += 1
            continue
        s = suggestions.get(rid)
        if s is None:
            counts["no_suggestion"] += 1
            continue
        mine_sent = str(ws.cell(r, col["gold_sentiment"]).value or "").strip().lower()
        mine_asp = aspect_set(ws.cell(r, col["gold_aspects"]).value)
        same = mine_sent == s["suggested_sentiment"] and mine_asp == aspect_set(s["suggested_aspects"])
        src.value = "claude_checked" if same else "claude_edited"
        counts[src.value] += 1
        if not same:
            diffs.append((rid, ws.cell(r, col["split"]).value, mine_sent, ";".join(sorted(mine_asp)),
                          s["suggested_sentiment"], ";".join(sorted(aspect_set(s["suggested_aspects"])))))

    letter = ws.cell(1, col["label_source"]).column_letter
    dv = DataValidation(type="list", formula1='"' + ",".join(CHOICES) + '"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f"{letter}2:{letter}{ws.max_row}")
    ws.column_dimensions[letter].width = 16
    wb.save(xlsx)
    return {"counts": counts, "diffs": diffs}


def main() -> int:
    res = mark()
    c = res["counts"]
    print(f"claude_checked (same as suggestion): {c['claude_checked']}")
    print(f"claude_edited  (you changed it)    : {c['claude_edited']}")
    print(f"already marked, left alone         : {c['kept']}")
    if c["no_suggestion"]:
        print(f"rows with no suggestion (left empty = human): {c['no_suggestion']}")
    if res["diffs"]:
        print("\nrows where your label differs from Claude's suggestion:")
        print(f"  {'id':6} {'split':5} {'yours':40} claude")
        for rid, split, ms, ma, cs, ca in res["diffs"]:
            print(f"  {rid:6} {split:5} {ms + ' / ' + ma:40} {cs} / {ca}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
