"""Offline tests for eval-set selection and gold-label validation."""
import csv
from collections import Counter

from app.taxonomy import ASPECTS, SENTIMENTS
from scripts.build_gold import validate
from scripts.select_eval_set import select


def load_reviews():
    with open("data/reviews.csv", encoding="utf-8", newline="") as f:
        return list(csv.DictReader(f))


def test_taxonomy_is_fixed():
    assert SENTIMENTS == ("positive", "neutral", "negative")
    assert ASPECTS == ("taste", "price", "service", "comfort", "portion", "other")


def test_selection_is_deterministic_balanced_and_disjoint():
    rows = load_reviews()
    a, b = select(rows), select(rows)
    assert [r["review_id"] for r in a] == [r["review_id"] for r in b]
    by_split = Counter(r["split"] for r in a)
    assert by_split == {"dev": 10, "test": 30}
    assert len({r["review_id"] for r in a}) == 40
    for split, per in (("dev", 2), ("test", 6)):
        assert Counter(r["rating"] for r in a if r["split"] == split) == {s: per for s in "12345"}


def _row(rid, text, sentiment="positive", aspects="taste", evidence=""):
    return {"review_id": rid, "split": "dev", "review_text": text, "gold_sentiment": sentiment,
            "gold_aspects": aspects, "gold_evidence": evidence, "notes": ""}


def test_validate_accepts_good_row():
    reviews = {"w001": "อร่อยมาก แต่แพงไปหน่อย"}
    gold, errors = validate([_row("w001", reviews["w001"], "positive", "price; taste", "อร่อยมาก | แพงไปหน่อย")], reviews)
    assert errors == []
    assert gold[0]["gold_aspects"] == ["taste", "price"]
    assert gold[0]["gold_evidence"] == ["อร่อยมาก", "แพงไปหน่อย"]


def test_validate_reports_every_problem():
    reviews = {"w001": "อร่อยมาก", "w002": "ร้านสะอาด"}
    rows = [
        _row("w001", "อร่อยมาก", "mixed", "taste;delivery", "ไม่อร่อย"),
        _row("w002", "ร้านสะอาด (edited)"),
        _row("w999", "x"),
    ]
    _, errors = validate(rows, reviews)
    text = "\n".join(errors)
    assert "gold_sentiment 'mixed'" in text
    assert "unknown aspect(s) ['delivery']" in text
    assert "evidence not found verbatim" in text
    assert "review_text was edited" in text
    assert "w999" in text and "not found" in text


def test_empty_aspects_allowed():
    reviews = {"w001": "ไปร้านนี้มาเมื่อวาน"}
    gold, errors = validate([_row("w001", reviews["w001"], "neutral", "")], reviews)
    assert errors == [] and gold[0]["gold_aspects"] == []


def test_label_workbook_round_trip(tmp_path, monkeypatch):
    """make_label_sheet -> (simulated human edit) -> build_gold reader gets the same text back."""
    from openpyxl import load_workbook

    from scripts import build_gold, make_label_sheet

    reviews = {r["review_id"]: r for r in load_reviews()}
    rows = [{"review_id": "w001", "split": "dev", "review_text": reviews["w001"]["review_text"]}]
    path = tmp_path / "labeling.xlsx"
    make_label_sheet.build(rows, path)

    wb = load_workbook(path)
    ws = wb["labels"]
    ws["D2"], ws["E2"] = "positive", "taste"
    wb.save(path)

    monkeypatch.setattr(build_gold, "XLSX", path)
    got, source = build_gold.read_label_rows()
    assert source == path
    gold, errors = validate(got, {k: v["review_text"] for k, v in reviews.items()})
    assert errors == []
    assert gold[0]["gold_sentiment"] == "positive" and gold[0]["gold_aspects"] == ["taste"]


def test_guide_refresh_keeps_labels_and_dropdown(tmp_path):
    from openpyxl import load_workbook

    from scripts import make_label_sheet
    from scripts.label_guide_th import ASPECTS_TH

    path = tmp_path / "labeling.xlsx"
    make_label_sheet.build([{"review_id": "w001", "split": "dev", "review_text": "อร่อยมาก"}], path)
    wb = load_workbook(path)
    wb["labels"]["D2"], wb["labels"]["E2"] = "positive", "taste"
    wb.save(path)

    make_label_sheet.refresh_guide(path)
    wb = load_workbook(path)
    assert wb["labels"]["D2"].value == "positive" and wb["labels"]["E2"].value == "taste"
    assert wb["labels"].data_validations.dataValidation  # dropdown still there
    guide_text = " ".join(str(c.value) for row in wb["guide"].iter_rows() for c in row if c.value)
    assert all(a in guide_text for a, *_ in ASPECTS_TH)
    assert [a for a, *_ in ASPECTS_TH] == list(ASPECTS)


def test_ai_drafts_must_be_reviewed():
    reviews = {"w001": "อร่อยมาก"}
    row = _row("w001", "อร่อยมาก")
    _, errors = validate([{**row, "label_source": "claude_draft"}], reviews)
    assert any("not reviewed yet" in e for e in errors)
    gold, errors = validate([{**row, "label_source": "claude_edited"}], reviews)
    assert errors == [] and gold[0]["label_source"] == "claude_edited"
    gold, _ = validate([row], reviews)
    assert gold[0]["label_source"] == "human"


def test_mark_label_source_compares_with_suggestions(tmp_path):
    from openpyxl import load_workbook

    from scripts import make_label_sheet, mark_label_source

    path = tmp_path / "l.xlsx"
    rows = [{"review_id": i, "split": "test", "review_text": "x"} for i in ("w001", "w002", "w003")]
    make_label_sheet.build(rows, path)
    wb = load_workbook(path)
    ws = wb["labels"]
    ws["D2"], ws["E2"] = "positive", "price;taste"      # same set as suggestion, other order
    ws["D3"], ws["E3"] = "neutral", "taste"             # different sentiment
    ws["D4"], ws["E4"] = "negative", "taste"
    ws["H1"], ws["H4"] = "label_source", "human"        # already marked: must be kept
    wb.save(path)
    sugg = {k: {"suggested_sentiment": s, "suggested_aspects": a}
            for k, s, a in (("w001", "positive", "taste;price"), ("w002", "negative", "taste"),
                            ("w003", "positive", "taste"))}
    res = mark_label_source.mark(path, sugg)
    ws = load_workbook(path)["labels"]
    assert [ws[f"H{r}"].value for r in (2, 3, 4)] == ["claude_checked", "claude_edited", "human"]
    assert res["counts"]["kept"] == 1 and [d[0] for d in res["diffs"]] == ["w002"]
